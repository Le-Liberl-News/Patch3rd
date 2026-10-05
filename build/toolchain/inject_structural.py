from __future__ import annotations

import argparse
import difflib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


TEXT_INSTRUCTION = re.compile(
    r"^(?P<indent>\s*)(?:TextTalk|TextTalkNamed|TextMessage|ScMenuSetTitle)\b.*\{\s*$"
)
MENU_INSTRUCTION = re.compile(r"^(?P<indent>\s*)Menu\b")
MENU_ITEM = re.compile(r'^(?P<prefix>\s*)"(?P<text>(?:\\.|[^"\\])*)"(?P<suffix>\s*//.*)?$')
MENU_ADD = re.compile(
    r'^(?P<prefix>\s*ED6MenuAdd\s+\S+\s+")(?P<text>(?:\\.|[^"\\])*)(?P<suffix>".*)$'
)
HASH_CONTROL_SUFFIX = r"(?:[A-Z]|v|\|)"
LEADING_TAGS = re.compile(
    r"^(?:(?:#[0-9]+" + HASH_CONTROL_SUFFIX + r")|(?:\{(?:color|0x)[^}]*\}))*"
)
HIDDEN_TAG = re.compile(r"#[0-9]+" + HASH_CONTROL_SUFFIX)
# Evolution voice tags (#<id>v): never shown nor translated.
VOICE_TAG = re.compile(r"#[0-9]+v")
HIDDEN_TAG_FOR_COMPARE = re.compile(r"#[0-9]+" + HASH_CONTROL_SUFFIX)
ENGINE_CONTROL = re.compile(r"#[0-9]+" + HASH_CONTROL_SUFFIX + r"|\{[^{}\r\n]+\}")
WAIT_MARKER = "***wait***"
CHARACTER_NAME = re.compile(r'^(?P<prefix>\s+name\s+")(?P<name>(?:\\.|[^"\\])*)(?P<suffix>"\s*)$')
SET_NAME = re.compile(r'^(?P<prefix>\s*TextSetName\s+")(?P<name>(?:\\.|[^"\\])*)(?P<suffix>".*)$')
NAMED_TALK = re.compile(
    r'^(?P<prefix>\s*TextTalkNamed\s+\S+\s+")(?P<name>(?:\\.|[^"\\])*)(?P<suffix>".*)$'
)
FRENCH_GLYPH_ENCODING = str.maketrans({
    "À": "ﾃ", "Â": "｡", "Ä": "､", "Ç": "ｦ", "È": "ｧ", "É": "ｨ",
    "Ê": "ｩ", "Ë": "ｪ", "Î": "ｫ", "Ï": "ｬ", "Ô": "ｭ", "Ö": "ﾊ",
    "Ù": "ｮ", "Û": "ﾂ", "Ü": "ｰ", "Ÿ": "ﾂ", "Œ": "ﾀ", "à": "ｱ",
    "â": "ｲ", "ä": "ｳ", "ç": "ｴ", "è": "ｵ", "é": "ｶ", "ê": "ｷ",
    "ë": "ｸ", "î": "ｹ", "ï": "ｺ", "ô": "ｻ", "ö": "ﾉ", "ù": "ｼ",
    "û": "ｽ", "ü": "ｾ", "ÿ": "ｿ", "œ": "ﾁ", "°": "ﾄ", "«": "ﾅ",
    "»": "ﾆ", "\N{NO-BREAK SPACE}": "ﾈ", "\N{NARROW NO-BREAK SPACE}": "ﾈ",
    "…": "...", "\N{HAIR SPACE}": "ﾋ", "’": "'", "♡": "㈱", "♥": "㈱",
    "❤": "㈱", "–": "-", "—": "-",
})


@dataclass
class Slot:
    kind: str
    start: int
    end: int
    english: str
    indent: str = ""
    prefix: str = ""
    suffix: str = ""
    raw: str = ""


@dataclass(frozen=True)
class NameRow:
    row: int
    english: str
    translation: str


@dataclass(frozen=True)
class WorkbookNames:
    characters: tuple[NameRow, ...]
    punctual: dict[str, str]


def clean_dialogue(raw: str) -> str:
    value = raw
    # Voices are never shown or translated; injection puts them back at the
    # start of the matching French page.
    value = VOICE_TAG.sub("", value)
    while re.match(r"^#[0-9]+" + HASH_CONTROL_SUFFIX, value):
        value = re.sub(r"^#[0-9]+" + HASH_CONTROL_SUFFIX, "", value, count=1)
    return value.replace("{wait}", "").rstrip("\n")


def unescape_quoted(value: str) -> str:
    return value.replace(r'\"', '"').replace(r"\\", "\\")


def escape_quoted(value: str) -> str:
    return value.replace("\\", r"\\").replace('"', r'\"')


def scan_slots(lines: list[str], include_empty: bool = False) -> list[Slot]:
    slots: list[Slot] = []
    index = 0
    while index < len(lines):
        bare = lines[index].rstrip("\r\n")
        text_match = TEXT_INSTRUCTION.match(bare)
        if text_match:
            block_indent = text_match.group("indent")
            page_start = index + 1
            cursor = page_start
            while cursor < len(lines):
                marker = lines[cursor].rstrip("\r\n")
                if marker == block_indent + "} {" or marker == block_indent + "}":
                    content = lines[page_start:cursor]
                    if not content:
                        raise ValueError(f"Empty dialogue page at line {index + 1}")
                    # Indentation of the first non-blank line: a page may start with a blank one.
                    content_indent = re.match(r"^\s*", next((line for line in content if line.strip()), content[0])).group(0)
                    raw = "\n".join(
                        line.rstrip("\r\n")[len(content_indent):] for line in content
                    )
                    prefix = LEADING_TAGS.match(raw).group(0)
                    suffix = "{wait}" if raw.endswith("{wait}") else ""
                    slots.append(
                        Slot(
                            "text", page_start, cursor, clean_dialogue(raw),
                            content_indent, prefix, suffix, raw,
                        )
                    )
                    if marker.endswith("} {"):
                        page_start = cursor + 1
                        cursor += 1
                        continue
                    index = cursor + 1
                    break
                cursor += 1
            else:
                raise ValueError(f"Unclosed dialogue at line {index + 1}")
            continue

        menu_match = MENU_INSTRUCTION.match(bare)
        if menu_match:
            base_indent = len(menu_match.group("indent"))
            cursor = index + 1
            while cursor < len(lines):
                candidate = lines[cursor].rstrip("\r\n")
                indent = len(re.match(r"^\s*", candidate).group(0))
                if candidate.strip() and indent <= base_indent:
                    break
                item = MENU_ITEM.match(candidate)
                if item:
                    slots.append(
                        Slot(
                            "quoted",
                            cursor,
                            cursor + 1,
                            unescape_quoted(item.group("text")),
                            prefix=item.group("prefix") + '"',
                            suffix='"' + (item.group("suffix") or ""),
                        )
                    )
                cursor += 1
            index = cursor
            continue

        menu_add = MENU_ADD.match(bare)
        if menu_add:
            slots.append(
                Slot(
                    "quoted",
                    index,
                    index + 1,
                    unescape_quoted(menu_add.group("text")),
                    prefix=menu_add.group("prefix"),
                    suffix=menu_add.group("suffix"),
                )
            )
        index += 1
    # The sheet writer and the historical injector ignore empty English
    # operands (for example a page containing only {wait}). They must not
    # consume a translated worksheet row.
    return slots if include_empty else [slot for slot in slots if slot.english != ""]


def workbook_rows(path: Path) -> list[tuple[int, str, str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = []
    original_seen = False
    for row, values in enumerate(sheet.iter_rows(min_col=3, max_col=5, values_only=True), 1):
        japanese, english, translation = values
        if japanese == "ORIGINAL":
            original_seen = True
            continue
        if not original_seen:
            continue
        if english not in (None, ""):
            rows.append((row, str(english).rstrip("\n"), "" if translation is None else str(translation)))
    workbook.close()
    return rows


def workbook_names(path: Path) -> WorkbookNames:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    characters = []
    punctual = {}
    in_punctual = False
    for row, (japanese, english, translation) in enumerate(sheet.iter_rows(
        min_row=6, min_col=3, max_col=5, values_only=True
    ), 6):
        if japanese == "ORIGINAL":
            break
        if japanese == "Noms ponctuels":
            in_punctual = True
            continue
        if in_punctual:
            if english in (None, ""):
                continue
            entry = NameRow(row, str(english), "" if translation is None else str(translation))
            previous = punctual.get(entry.english)
            if previous is not None and previous != entry.translation:
                raise ValueError(
                    f"{path.name}: traductions ponctuelles contradictoires pour {entry.english!r}"
                )
            punctual[entry.english] = entry.translation
        else:
            # Every row between the fixed name-table header and the next
            # section represents one npc slot, including deliberately
            # anonymous characters whose JP/EN/FR cells are all empty.
            entry = NameRow(
                row,
                "" if english is None else str(english),
                "" if translation is None else str(translation),
            )
            characters.append(entry)
    workbook.close()
    return WorkbookNames(tuple(characters), punctual)


def apply_workbook_names(lines: list[str], names: WorkbookNames, label: str) -> int:
    character_slots = []
    inside_npc = False
    for index, line in enumerate(lines):
        bare = line.rstrip("\r\n")
        if re.match(r"^npc char\[\d+\]:\s*$", bare):
            inside_npc = True
            continue
        if inside_npc and bare and not bare[0].isspace():
            inside_npc = False
        if inside_npc:
            match = CHARACTER_NAME.match(bare)
            if match:
                character_slots.append((index, match))

    # Included scripts (_1, _2, ...) inherit the parent's runtime character
    # table and therefore legitimately repeat its XLSX rows without declaring
    # local npc blocks. Their names are injected once in the owning CLM.
    aligned_slots = []
    if character_slots:
        # An anonymous slot (empty or blank name) has nothing to translate and
        # may legitimately be absent from the workbook's name table.
        named_slots = [
            (index, match)
            for index, match in character_slots
            if unescape_quoted(match.group("name")).strip()
        ]
        if not names.characters and named_slots:
            raise ValueError(
                f"{label}: 0 noms de personnages XLSX mais "
                f"{len(named_slots)} slots CLM"
            )
        slot_names = [unescape_quoted(match.group("name")) for _, match in character_slots]

        forward = []
        cursor = 0
        for entry in names.characters:
            while cursor < len(slot_names) and slot_names[cursor] != entry.english:
                cursor += 1
            if cursor == len(slot_names):
                raise ValueError(
                    f"{label}: impossible d'aligner le nom XLSX ligne {entry.row} "
                    f"{entry.english!r} sur les slots CLM"
                )
            forward.append(cursor)
            cursor += 1

        backward = []
        cursor = len(slot_names) - 1
        for entry in reversed(names.characters):
            while cursor >= 0 and slot_names[cursor] != entry.english:
                cursor -= 1
            if cursor < 0:
                raise ValueError(
                    f"{label}: impossible d'aligner le nom XLSX ligne {entry.row} "
                    f"{entry.english!r} sur les slots CLM"
                )
            backward.append(cursor)
            cursor -= 1
        backward.reverse()

        if forward != backward:
            first = next(i for i, pair in enumerate(zip(forward, backward)) if pair[0] != pair[1])
            entry = names.characters[first]
            raise ValueError(
                f"{label}: alignement ambigu du nom XLSX ligne {entry.row} "
                f"{entry.english!r} (slots CLM {forward[first] + 1} ou {backward[first] + 1})"
            )
        aligned_slots = [character_slots[index] for index in forward]

    changed = 0
    for entry, (index, match) in zip(names.characters, aligned_slots):
        original = unescape_quoted(match.group("name"))
        if original != entry.english:
            raise ValueError(
                f"{label}: nom ligne {entry.row} ne correspond pas au slot CLM: "
                f"{entry.english!r} != {original!r}"
            )
        if entry.translation and entry.translation != original:
            newline = "\r\n" if lines[index].endswith("\r\n") else "\n"
            replacement = encode_french_glyphs(entry.translation)
            lines[index] = (
                match.group("prefix") + escape_quoted(replacement) + match.group("suffix") + newline
            )
            changed += 1

    for index, line in enumerate(lines):
        bare = line.rstrip("\r\n")
        match = SET_NAME.match(bare) or NAMED_TALK.match(bare)
        if not match:
            continue
        original = unescape_quoted(match.group("name"))
        replacement = names.punctual.get(original, "")
        if replacement and replacement != original:
            newline = "\r\n" if line.endswith("\r\n") else "\n"
            replacement = encode_french_glyphs(replacement)
            lines[index] = (
                match.group("prefix") + escape_quoted(replacement) + match.group("suffix") + newline
            )
            changed += 1
    return changed


def normalized(value: str) -> str:
    value = value.replace("\r\n", "\n")
    value = value.replace("{wait}", "").replace(WAIT_MARKER, "")
    value = HIDDEN_TAG_FOR_COMPARE.sub("", value)
    return value.rstrip("\n")


def encode_french_glyphs(value: str) -> str:
    """Map French glyphs to the custom CP932 slots used by the SC patch font.

    Text is composed first (NFC): an "É" typed as E + U+0301, as some editors
    and Mac keyboards produce, would otherwise have no glyph.
    """
    return unicodedata.normalize("NFC", value).translate(FRENCH_GLYPH_ENCODING)


def _control_occurrences(value: str) -> list[re.Match[str]]:
    return list(ENGINE_CONTROL.finditer(value))


def _hidden_control_indexes(source: str, displayed_english: str) -> set[int]:
    source_controls = _control_occurrences(source)
    displayed_controls = ENGINE_CONTROL.findall(displayed_english)
    matcher = difflib.SequenceMatcher(
        a=[match.group(0) for match in source_controls],
        b=displayed_controls,
        autojunk=False,
    )
    visible: set[int] = set()
    for block in matcher.get_matching_blocks():
        visible.update(range(block.a, block.a + block.size))
    return set(range(len(source_controls))) - visible


def _control_family(token: str) -> str:
    match = re.fullmatch(r"#[0-9]+([A-Z]|v)", token)
    return match.group(1) if match else token


def _translation_supplies(token: str, translation: str) -> bool:
    if token in translation:
        return True
    color = re.fullmatch(r"#([0-9]+)C", token)
    if color and re.search(r"\{color\s+" + re.escape(color.group(1)) + r"\}", translation):
        return True
    # Layout/timing controls are legitimately rewritten for French line lengths.
    family = _control_family(token)
    return family in {"W", "S", "A", "P"} and any(
        _control_family(candidate) == family for candidate in ENGINE_CONTROL.findall(translation)
    )


def _insert_after_wait(value: str, ordinal: int, controls: str, next_line: bool = False) -> str:
    matches = list(re.finditer(r"\{wait\}", value))
    if ordinal <= 0 or ordinal > len(matches):
        raise ValueError(
            f"hidden controls {controls!r} follow wait #{ordinal}, but the translation "
            f"contains only {len(matches)} internal wait marker(s)"
        )
    position = matches[ordinal - 1].end()
    # Controls that open the next source line (voices, mostly) open the
    # next French line too.
    if next_line and value[position:position + 1] == "\n":
        position += 1
    return value[:position] + controls + value[position:]


def restore_hidden_controls(
    raw: str, displayed_english: str, translation: str, dropped: list[str] | None = None,
    relocate_ambiguous_line_controls: bool = False,
) -> str:
    """Restore only controls proven to have been hidden by the XLSX writer.

    Without ``dropped``, an ambiguous control position is an error. With it,
    such controls are omitted and reported in ``dropped``: a translation that
    leaves out internal pauses stays compilable and is caught in review.
    """
    source = raw.replace("\r\n", "\n")
    translated = translation.replace("\r\n", "\n")
    translated = translated.replace(WAIT_MARKER, "{wait}")
    occurrences = _control_occurrences(source)
    hidden = _hidden_control_indexes(source, displayed_english.replace(WAIT_MARKER, "{wait}"))

    leading: list[str] = []
    after_wait: dict[int, list[str]] = {}
    after_wait_line: set[int] = set()
    line_controls: dict[int, list[str]] = {}
    unresolved: list[str] = []
    for index, match in enumerate(occurrences):
        if index not in hidden:
            continue
        token, start, end = match.group(0), match.start(), match.end()
        before, after = source[:start], source[end:]
        if token == "{wait}" and not after.strip():
            if not translated.endswith("{wait}"):
                translated += "{wait}"
            continue
        if start == 0 or re.fullmatch(r"(?:#[0-9]+(?:[A-Z]|v))*", before):
            leading.append(token)
            continue
        wait_prefix = re.search(r"\{wait\}\s*(?:#[0-9]+(?:[A-Z]|v))*$", before)
        if wait_prefix:
            if not _translation_supplies(token, translated):
                ordinal = before[: wait_prefix.start()].count("{wait}") + 1
                after_wait.setdefault(ordinal, []).append(token)
                if "\n" in wait_prefix.group(0):
                    after_wait_line.add(ordinal)
            continue
        if _translation_supplies(token, translated):
            continue
        if before.endswith("\n"):
            line_controls.setdefault(before.count("\n"), []).append(token)
            continue
        unresolved.append(f"{token} at source offset {start}")

    missing_leading = [token for token in leading if not _translation_supplies(token, translated)]
    if missing_leading:
        translated = "".join(missing_leading) + translated
    for ordinal in sorted(after_wait, reverse=True):
        controls = "".join(after_wait[ordinal])
        try:
            translated = _insert_after_wait(
                translated, ordinal, controls, ordinal in after_wait_line
            )
        except ValueError:
            if dropped is None:
                raise
            unresolved.append(f"{controls} after source wait #{ordinal}")
    if line_controls:
        lines = translated.split("\n")
        source_lines = source.split("\n")
        if len(lines) != len(source_lines) and relocate_ambiguous_line_controls:
            # PatchSC: controls of a reflowed page open the French page.
            relocated = [
                token
                for _, tokens in sorted(line_controls.items())
                for token in tokens
                if not _translation_supplies(token, translated)
            ]
            translated = "".join(relocated) + translated
        elif len(lines) != len(source_lines):
            unresolved.extend(
                f"{''.join(tokens)} at start of source line {line + 1}"
                for line, tokens in sorted(line_controls.items())
            )
        else:
            for line, tokens in sorted(line_controls.items()):
                missing = [
                    token for token in tokens if not _translation_supplies(token, lines[line])
                ]
                if missing:
                    lines[line] = "".join(missing) + lines[line]
            translated = "\n".join(lines)
    if unresolved and dropped is not None:
        dropped.extend(unresolved)
    elif unresolved:
        raise ValueError(
            "hidden control position is ambiguous after translation: " + "; ".join(unresolved)
        )
    return translated


def replace_slot(
    lines: list[str], slot: Slot, displayed_english: str, translation: str,
    dropped: list[str] | None = None, relocate_ambiguous_line_controls: bool = False,
) -> None:
    translation = encode_french_glyphs(translation)
    newline = "\r\n" if lines[slot.start].endswith("\r\n") else "\n"
    if slot.kind == "quoted":
        lines[slot.start:slot.end] = [slot.prefix + escape_quoted(translation) + slot.suffix + newline]
        return
    translated_lines = restore_hidden_controls(
        slot.raw, displayed_english, translation, dropped, relocate_ambiguous_line_controls
    ).split("\n")
    lines[slot.start:slot.end] = [slot.indent + value + newline for value in translated_lines]


def inject_file(workbook: Path, base: Path, output: Path) -> dict[str, object]:
    lines = base.read_text(encoding="utf-8").splitlines(keepends=True)
    slots = scan_slots(lines)
    rows = workbook_rows(workbook)
    if len(rows) != len(slots):
        raise ValueError(f"{workbook.name}: {len(rows)} XLSX rows but {len(slots)} CLM slots")
    mismatches = []
    for position, ((row, english, _), slot) in enumerate(zip(rows, slots)):
        if normalized(english) != normalized(slot.english):
            mismatches.append(
                f"row {row}/slot {position + 1}: {english[:50]!r} != {slot.english[:50]!r}"
            )
    if mismatches:
        raise ValueError(f"{workbook.name}: structural mismatch: " + "; ".join(mismatches[:3]))

    changed = 0
    for (_, english, translation), slot in reversed(list(zip(rows, slots))):
        if translation and normalized(translation) != normalized(english):
            replace_slot(lines, slot, english, translation)
            changed += 1
    names_changed = apply_workbook_names(lines, workbook_names(workbook), workbook.name)
    output.write_text("".join(lines), encoding="utf-8", newline="")
    return {
        "file": workbook.name,
        "slots": len(slots),
        "changed": changed,
        "names_changed": names_changed,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Structurally inject Sky translation workbooks")
    parser.add_argument("--scenario", type=Path, default=Path("scenario"))
    parser.add_argument("--base-clm", type=Path, default=Path("base_clm"))
    parser.add_argument("--output", type=Path, default=Path("translated_clm"))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    bases = {path.stem.strip().upper(): path for path in args.base_clm.glob("*.clm")}
    failures = []
    total_slots = total_changed = 0
    workbooks = sorted(args.scenario.glob("*.xlsx"))
    for index, workbook in enumerate(workbooks, 1):
        base = bases.get(workbook.stem.strip().upper())
        if base is None:
            failures.append(f"{workbook.name}: base CLM missing")
            continue
        try:
            result = inject_file(workbook, base, args.output / base.name)
            total_slots += int(result["slots"])
            total_changed += int(result["changed"])
        except Exception as error:
            message = str(error)
            failures.append(
                message if message.startswith(f"{workbook.name}:") else f"{workbook.name}: {message}"
            )
        if index % 50 == 0:
            print(f"{index}/{len(workbooks)}", flush=True)
    print(f"files={len(workbooks) - len(failures)} slots={total_slots} changed={total_changed}")
    for failure in failures:
        print("ERROR:", failure)
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
